from io import BytesIO
from datetime import datetime
from html import escape
from pathlib import Path

try:
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.lib.colors import HexColor
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


def _s(value) -> str:
    return "" if value is None else str(value)


def generate_txt_export(analysis_results: dict) -> str:
    out = "=" * 80 + "\n"
    out += "RESUME ANALYSIS REPORT\n"
    out += f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n"
    out += "=" * 80 + "\n\n"

    if analysis_results.get("resume_summary"):
        out += "RESUME SUMMARY\n" + "-" * 40 + "\n"
        out += _s(analysis_results["resume_summary"]) + "\n\n"

    if analysis_results.get("job_requirements"):
        out += "JOB REQUIREMENTS\n" + "-" * 40 + "\n"
        for req in analysis_results["job_requirements"]:
            out += f"- {_s(req)}\n"
        out += "\n"

    if analysis_results.get("matching_skills"):
        out += "MATCHING SKILLS\n" + "-" * 40 + "\n"
        for skill in analysis_results["matching_skills"]:
            if isinstance(skill, dict):
                out += f"[+] {_s(skill.get('skill'))}\n"
                out += f"    Evidence: \"{_s(skill.get('evidence'))}\"\n\n"
            else:
                out += f"[+] {_s(skill)}\n\n"

    if analysis_results.get("missing_areas"):
        out += "MISSING OR WEAK AREAS\n" + "-" * 40 + "\n"
        for area in analysis_results["missing_areas"]:
            out += f"[!] {_s(area)}\n"
        out += "\n"

    if analysis_results.get("suggestions"):
        out += "IMPROVEMENT SUGGESTIONS\n" + "-" * 40 + "\n"
        for i, sugg in enumerate(analysis_results["suggestions"], 1):
            if isinstance(sugg, dict):
                out += f"\n{i}. Original:\n   \"{_s(sugg.get('original'))}\"\n"
                out += f"\n   Improved:\n   \"{_s(sugg.get('improved'))}\"\n"
                out += f"\n   Why: {_s(sugg.get('reason'))}\n"
            else:
                out += f"\n{i}. {_s(sugg)}\n"

    if analysis_results.get("improved_resume"):
        out += "\n" + "=" * 80 + "\n"
        out += "IMPROVED RESUME\n"
        out += "=" * 80 + "\n\n"
        out += _s(analysis_results["improved_resume"])

    return out


def generate_pdf_export(analysis_results: dict) -> bytes:
    if not PDF_AVAILABLE:
        raise ValueError("PDF export not available. Install reportlab: pip install reportlab")

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=0.5 * inch, bottomMargin=0.5 * inch)
    styles = getSampleStyleSheet()
    story = []

    title_style = ParagraphStyle(
        "CustomTitle", parent=styles["Heading1"], fontSize=24,
        textColor=HexColor("#667eea"), spaceAfter=12, alignment=1,
    )
    story.append(Paragraph("Resume Analysis Report", title_style))
    story.append(Spacer(1, 0.2 * inch))

    if analysis_results.get("resume_summary"):
        story.append(Paragraph("Resume Summary", styles["Heading2"]))
        story.append(Paragraph(escape(_s(analysis_results["resume_summary"])), styles["Normal"]))
        story.append(Spacer(1, 0.2 * inch))

    if analysis_results.get("job_requirements"):
        story.append(Paragraph("Job Requirements", styles["Heading2"]))
        for req in analysis_results["job_requirements"]:
            story.append(Paragraph(f"&bull; {escape(_s(req))}", styles["Normal"]))
        story.append(Spacer(1, 0.2 * inch))

    if analysis_results.get("matching_skills"):
        story.append(Paragraph("Matching Skills [Match]", styles["Heading2"]))
        for skill in analysis_results["matching_skills"]:
            if isinstance(skill, dict):
                story.append(Paragraph(f"<b>{escape(_s(skill.get('skill')))}</b>", styles["Normal"]))
                evidence = _s(skill.get("evidence"))
                if evidence:
                    story.append(Paragraph(f"<i>Evidence: \"{escape(evidence)}\"</i>", styles["Normal"]))
            else:
                story.append(Paragraph(f"<b>{escape(_s(skill))}</b>", styles["Normal"]))
            story.append(Spacer(1, 0.1 * inch))
        story.append(Spacer(1, 0.1 * inch))

    if analysis_results.get("missing_areas"):
        story.append(Paragraph("Missing or Weak Areas", styles["Heading2"]))
        for area in analysis_results["missing_areas"]:
            story.append(Paragraph(f"[!] {escape(_s(area))}", styles["Normal"]))
        story.append(Spacer(1, 0.2 * inch))

    if analysis_results.get("suggestions"):
        story.append(PageBreak())
        story.append(Paragraph("Improvement Suggestions", styles["Heading2"]))
        for i, sugg in enumerate(analysis_results["suggestions"], 1):
            story.append(Paragraph(f"<b>Suggestion {i}</b>", styles["Heading3"]))
            if isinstance(sugg, dict):
                story.append(Paragraph(f"<b>Original:</b> {escape(_s(sugg.get('original')))}", styles["Normal"]))
                story.append(Paragraph(f"<b>Improved:</b> {escape(_s(sugg.get('improved')))}", styles["Normal"]))
                story.append(Paragraph(f"<b>Why:</b> {escape(_s(sugg.get('reason')))}", styles["Normal"]))
            else:
                story.append(Paragraph(escape(_s(sugg)), styles["Normal"]))
            story.append(Spacer(1, 0.15 * inch))

    if analysis_results.get("improved_resume"):
        story.append(PageBreak())
        story.append(Paragraph("Improved Resume", styles["Heading1"]))
        story.append(Spacer(1, 0.2 * inch))
        improved_text = escape(_s(analysis_results["improved_resume"])).replace("\n", "<br/>")
        story.append(Paragraph(improved_text, styles["Normal"]))

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


def generate_docx_export(analysis_results: dict) -> bytes:
    if not DOCX_AVAILABLE:
        raise ValueError("DOCX export not available. Install python-docx: pip install python-docx")

    doc = Document()
    title = doc.add_heading("Resume Analysis Report", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if analysis_results.get("resume_summary"):
        doc.add_heading("Resume Summary", level=1)
        doc.add_paragraph(_s(analysis_results["resume_summary"]))

    if analysis_results.get("job_requirements"):
        doc.add_heading("Job Requirements", level=1)
        for req in analysis_results["job_requirements"]:
            doc.add_paragraph(_s(req), style="List Bullet")

    if analysis_results.get("matching_skills"):
        doc.add_heading("Matching Skills [Match]", level=1)
        for skill in analysis_results["matching_skills"]:
            p = doc.add_paragraph()
            if isinstance(skill, dict):
                p.add_run(_s(skill.get("skill"))).bold = True
                evidence = _s(skill.get("evidence"))
                if evidence:
                    doc.add_paragraph(f"Evidence: \"{evidence}\"", style="List Bullet 2")
            else:
                p.add_run(_s(skill)).bold = True

    if analysis_results.get("missing_areas"):
        doc.add_heading("Missing or Weak Areas", level=1)
        for area in analysis_results["missing_areas"]:
            doc.add_paragraph(f"[!] {_s(area)}", style="List Bullet")

    if analysis_results.get("suggestions"):
        doc.add_page_break()
        doc.add_heading("Improvement Suggestions", level=1)
        for i, sugg in enumerate(analysis_results["suggestions"], 1):
            doc.add_heading(f"Suggestion {i}", level=2)
            if isinstance(sugg, dict):
                for label, key in (("Original: ", "original"), ("Improved: ", "improved"), ("Why: ", "reason")):
                    p = doc.add_paragraph()
                    p.add_run(label).bold = True
                    p.add_run(_s(sugg.get(key)))
            else:
                doc.add_paragraph(_s(sugg))

    if analysis_results.get("improved_resume"):
        doc.add_page_break()
        doc.add_heading("Improved Resume", level=1)
        doc.add_paragraph(_s(analysis_results["improved_resume"]))

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


# ---------- Formatted resume export (A4, one page) ----------
import base64
import shutil
import subprocess
import tempfile

RESUME_FONTS = {
    "Arial": {"pdf": ("Helvetica", "Helvetica-Bold", "Helvetica-Oblique"), "docx": "Arial", "css": "Arial, Helvetica, sans-serif"},
    "Calibri": {"pdf": ("Helvetica", "Helvetica-Bold", "Helvetica-Oblique"), "docx": "Calibri", "css": "Calibri, Carlito, sans-serif"},
    "Georgia": {"pdf": ("Times-Roman", "Times-Bold", "Times-Italic"), "docx": "Georgia", "css": "Georgia, serif"},
    "Times New Roman": {"pdf": ("Times-Roman", "Times-Bold", "Times-Italic"), "docx": "Times New Roman", "css": "'Times New Roman', Times, serif"},
}
ACCENTS = {"black": "#000000", "navy": "#1F3A5F", "gray": "#444444"}
MARGIN_IN = 0.6
PDF_SIZES = [11, 10.5, 10, 9.5, 9]   # tried in order until the resume fits one page
DOCX_BODY_PT = 10.5
MIME_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


TEMPLATES = {
    "classic":   {"label": "Classic",   "align": "left",   "heading": "rule",   "skills": "pills"},
    "modern":    {"label": "Modern",    "align": "left",   "heading": "band",   "skills": "pills"},
    "minimal":   {"label": "Minimal",   "align": "left",   "heading": "plain",  "skills": "inline"},
    "executive": {"label": "Executive", "align": "center", "heading": "center", "skills": "inline"},
}


def decode_photo(data_url: str):
    """Return raw image bytes from a data: URL, or None."""
    if not data_url or not data_url.startswith("data:image/"):
        return None
    try:
        return base64.b64decode(data_url.split(",", 1)[1])
    except Exception:
        return None


def _contact_line(d):
    c = d.get("contact") or {}
    return " | ".join(_s(v) for v in (c.get("address"), c.get("phone"), c.get("email")) if v)


def _sections(d):
    """Sections in display order: (title, kind, payload)."""
    out = []
    if d.get("summary"):
        out.append(("Profile", "text", _s(d["summary"])))
    if d.get("experience"):
        out.append(("Work Experience", "entries", d["experience"]))
    if d.get("education"):
        out.append(("Education", "edu", d["education"]))
    if d.get("skills"):
        out.append(("Skills", "chips", [_s(x) for x in d["skills"]]))
    if d.get("languages"):
        out.append(("Languages", "chips", [_s(x) for x in d["languages"]]))
    if d.get("certifications"):
        out.append(("Certifications", "bullets", [_s(x) for x in d["certifications"]]))
    return out


def _entry(item, kind):
    """Return (title, meta line, bullets) for one job or education item."""
    if not isinstance(item, dict):
        return _s(item), "", []
    if kind == "entries":
        title = item.get("title")
        org = ", ".join(_s(x) for x in (item.get("company"), item.get("location")) if x)
        bullets = item.get("bullets") or []
    else:
        title = item.get("degree")
        org = ", ".join(_s(x) for x in (item.get("school"), item.get("location")) if x)
        bullets = []
    meta = " | ".join(x for x in (org, _s(item.get("dates"))) if x)
    return _s(title), meta, [_s(b) for b in bullets]


def _pdf_once(d, font_key, accent_hex, photo, size, tpl="classic"):
    from io import BytesIO
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import inch
    from reportlab.lib.colors import HexColor, white
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                    TableStyle, Image, HRFlowable)

    cfg = TEMPLATES.get(tpl, TEMPLATES["classic"])
    regular, bold, italic = RESUME_FONTS[font_key]["pdf"]
    accent = HexColor(accent_hex)
    gray = HexColor("#555555")
    margin = MARGIN_IN * inch
    width = A4[0] - 2 * margin
    small = size - 0.5
    jal = 1 if cfg["align"] == "center" else 0
    hal = 1 if cfg["heading"] == "center" else 0

    name_s = ParagraphStyle("name", fontName=bold, fontSize=size + 10, leading=size + 12, textColor=accent, alignment=jal)
    headline_s = ParagraphStyle("headline", fontName=bold, fontSize=small, leading=size + 1.5, textColor=gray, spaceBefore=4, alignment=jal)
    contact_s = ParagraphStyle("contact", fontName=regular, fontSize=small, leading=small * 1.4, textColor=gray, spaceBefore=3, alignment=jal)
    head_s = ParagraphStyle("head", fontName=bold, fontSize=size + 0.5, leading=size + 3, textColor=accent, alignment=hal)
    body_s = ParagraphStyle("body", fontName=regular, fontSize=size, leading=size * 1.35, spaceAfter=2)
    role_s = ParagraphStyle("role", fontName=bold, fontSize=size, leading=size * 1.3)
    date_s = ParagraphStyle("date", fontName=regular, fontSize=small, leading=small * 1.3, textColor=gray, alignment=TA_RIGHT)
    org_s = ParagraphStyle("org", fontName=italic, fontSize=small, leading=small * 1.3, textColor=gray)
    bullet_s = ParagraphStyle("bullet", parent=body_s, leftIndent=12, bulletIndent=2, spaceAfter=1.5)
    chip_s = ParagraphStyle("chip", fontName=regular, fontSize=small, leading=small * 1.35)
    pill_s = ParagraphStyle("pill", parent=chip_s, alignment=1)
    band_s = ParagraphStyle("band", fontName=bold, fontSize=size + 0.5, leading=size + 3, textColor=white)

    def rule(thickness, before=0, after=4):
        return HRFlowable(width="100%", thickness=thickness, color=accent, spaceBefore=before, spaceAfter=after)

    def heading(title):
        t = escape(title.upper())
        kind = cfg["heading"]
        if kind == "band":
            tbl = Table([[Paragraph(t, band_s)]], colWidths=[width])
            tbl.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), accent),
                                     ("LEFTPADDING", (0, 0), (-1, -1), 6),
                                     ("TOPPADDING", (0, 0), (-1, -1), 3),
                                     ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
            return [Spacer(1, 4), tbl, Spacer(1, 4)]
        if kind == "plain":
            return [Paragraph(t, head_s), Spacer(1, 2)]
        if kind == "center":
            return [rule(0.6, before=4, after=2), Paragraph(t, head_s), rule(0.6, after=4)]
        return [Paragraph(t, head_s), rule(0.5, after=4)]

    story = []
    headline = ""
    first_job = (d.get("experience") or [None])[0]
    first_edu = (d.get("education") or [None])[0]
    if isinstance(first_job, dict):
        headline = _s(first_job.get("title"))
    if not headline and isinstance(first_edu, dict):
        headline = _s(first_edu.get("degree"))

    left = [Paragraph(escape(_s(d.get("name")).upper()), name_s)]
    if headline:
        left.append(Paragraph(escape(headline.upper()), headline_s))
    left.append(Paragraph(escape(_contact_line(d)), contact_s))

    if photo and cfg["align"] == "center":
        pw = 1.1 * inch
        img = Image(BytesIO(photo), width=pw, height=pw)
        img.hAlign = "CENTER"
        story.append(img)
        story += left
    elif photo:
        pw = 1.1 * inch
        head = Table([[left, Image(BytesIO(photo), width=pw, height=pw)]],
                     colWidths=[width - pw - 10, pw + 10])
        head.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                                  ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                  ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
        story.append(head)
    else:
        story += left
    story.append(rule(2.2, before=8, after=0))
    story.append(rule(0.5, after=6))

    for title, kind, payload in _sections(d):
        story += heading(title)

        if kind == "text":
            story.append(Paragraph(escape(payload), body_s))

        elif kind == "bullets":
            for b in payload:
                story.append(Paragraph(escape(b), bullet_s, bulletText="•"))

        elif kind == "chips":
            if cfg["skills"] == "inline":
                story.append(Paragraph(escape("  •  ".join(payload)), body_s))
            else:
                cols = 3
                cells = []
                for x in payload:
                    pill = Table([[Paragraph(escape(x), pill_s)]], colWidths=[width / cols - 8])
                    pill.setStyle(TableStyle([
                        ("BOX", (0, 0), (-1, -1), 0.7, accent),
                        ("TOPPADDING", (0, 0), (-1, -1), 2.5),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                        ("LEFTPADDING", (0, 0), (-1, -1), 4),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                    ]))
                    cells.append(pill)
                while len(cells) % cols:
                    cells.append("")
                grid = Table([cells[i:i + cols] for i in range(0, len(cells), cols)],
                             colWidths=[width / cols] * cols)
                grid.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                                          ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                          ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                                          ("TOPPADDING", (0, 0), (-1, -1), 2),
                                          ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
                story.append(grid)

        else:
            for item in payload:
                if not isinstance(item, dict):
                    story.append(Paragraph(escape(_s(item)), body_s))
                    continue
                if kind == "entries":
                    role = _s(item.get("title"))
                    org = ", ".join(_s(x) for x in (item.get("company"), item.get("location")) if x)
                    bullets = item.get("bullets") or []
                else:
                    role = _s(item.get("degree"))
                    org = ", ".join(_s(x) for x in (item.get("school"), item.get("location")) if x)
                    bullets = []

                row = Table([[Paragraph(escape(role), role_s),
                              Paragraph(escape(_s(item.get("dates"))), date_s)]],
                            colWidths=[width * 0.72, width * 0.28])
                row.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                                         ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                         ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                                         ("TOPPADDING", (0, 0), (-1, -1), 0),
                                         ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]))
                story.append(row)
                if org:
                    story.append(Paragraph(escape(org), org_s))
                for b in bullets:
                    story.append(Paragraph(escape(_s(b)), bullet_s, bulletText="•"))
                story.append(Spacer(1, 5))

    later = {"count": 0}

    def _count_later(canvas, document):
        later["count"] += 1

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=margin, rightMargin=margin,
                            topMargin=margin, bottomMargin=margin)
    doc.build(story, onFirstPage=lambda c, dd: None, onLaterPages=_count_later)
    return buf.getvalue(), later["count"] == 0


def build_resume_pdf(d, font_key, accent_hex, photo, tpl="classic"):
    pdf = b""
    for size in PDF_SIZES:
        pdf, fits = _pdf_once(d, font_key, accent_hex, photo, size, tpl)
        if fits:
            break
    return pdf


def build_resume_docx(d, font_key, accent_hex, photo):
    from io import BytesIO
    from docx import Document
    from docx.shared import Pt, Inches, Mm, RGBColor
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    hex6 = accent_hex.lstrip("#")
    accent = RGBColor(int(hex6[0:2], 16), int(hex6[2:4], 16), int(hex6[4:6], 16))
    gray = RGBColor(0x55, 0x55, 0x55)
    size = DOCX_BODY_PT
    text_w = 8.27 - 2 * MARGIN_IN   # A4 width in inches minus margins

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Inches(MARGIN_IN)
    sec.top_margin = sec.bottom_margin = Inches(MARGIN_IN)

    normal = doc.styles["Normal"]
    normal.font.name = RESUME_FONTS[font_key]["docx"]
    normal.font.size = Pt(size)
    normal.paragraph_format.space_after = Pt(2)

    def bottom_rule(p, sz):
        pPr = p._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        b = OxmlElement("w:bottom")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(sz))
        b.set(qn("w:space"), "1")
        b.set(qn("w:color"), hex6.upper())
        pbdr.append(b)
        pPr.append(pbdr)

    def para(text="", bold=False, italic=False, color=None, delta=0, before=0, after=2, rule=0):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        if text:
            r = p.add_run(text)
            r.bold, r.italic = bold, italic
            r.font.size = Pt(size + delta)
            if color is not None:
                r.font.color.rgb = color
        if rule:
            bottom_rule(p, rule)
        return p

    name = _s(d.get("name"))
    contact = _contact_line(d)
    if photo:
        table = doc.add_table(rows=1, cols=2)
        table.autofit = False
        left, right = table.rows[0].cells
        left.width, right.width = Inches(text_w - 1.2), Inches(1.2)
        r = left.paragraphs[0].add_run(name)
        r.bold = True
        r.font.size = Pt(size + 10)
        r.font.color.rgb = accent
        cr = left.add_paragraph().add_run(contact)
        cr.font.size = Pt(size - 0.5)
        cr.font.color.rgb = gray
        right.paragraphs[0].add_run().add_picture(BytesIO(photo), width=Inches(1.0))
    else:
        para(name, bold=True, color=accent, delta=10, after=2)
        para(contact, color=gray, delta=-0.5, after=4)
    para(before=4, after=8, rule=12)

    for title, kind, payload in _sections(d):
        para(title.upper(), bold=True, color=accent, delta=0.5, before=8, after=4, rule=6)

        if kind == "text":
            para(payload)
        elif kind == "chips":
            para("   •   ".join(payload))
        elif kind == "bullets":
            for b in payload:
                doc.add_paragraph(b, style="List Bullet")
        else:
            for item in payload:
                if not isinstance(item, dict):
                    para(_s(item))
                    continue
                if kind == "entries":
                    role = _s(item.get("title"))
                    org = ", ".join(_s(x) for x in (item.get("company"), item.get("location")) if x)
                    bullets = item.get("bullets") or []
                else:
                    role = _s(item.get("degree"))
                    org = ", ".join(_s(x) for x in (item.get("school"), item.get("location")) if x)
                    bullets = []

                p = para(before=3, after=0)
                p.paragraph_format.tab_stops.add_tab_stop(Inches(text_w), WD_TAB_ALIGNMENT.RIGHT)
                rr = p.add_run(role)
                rr.bold = True
                dr = p.add_run("\t" + _s(item.get("dates")))
                dr.font.color.rgb = gray
                dr.font.size = Pt(size - 0.5)
                if org:
                    para(org, italic=True, color=gray, delta=-0.5, after=1)
                for b in bullets:
                    doc.add_paragraph(_s(b), style="List Bullet")

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_resume_doc(docx_bytes):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise ValueError("DOC export needs LibreOffice installed. Download DOCX instead.")
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "formatted_resume.docx"
        src.write_bytes(docx_bytes)
        try:
            subprocess.run([soffice, "--headless", "--convert-to", "doc", "--outdir", tmp, str(src)],
                           check=True, capture_output=True, timeout=120)
        except (subprocess.SubprocessError, OSError) as e:
            raise ValueError(f"Could not convert to DOC: {e}")
        return (Path(tmp) / "formatted_resume.doc").read_bytes()


def build_resume_txt(d):
    lines = [_s(d.get("name")).upper()]
    contact = _contact_line(d)
    if contact:
        lines.append(contact)
    lines.append("")
    for title, kind, payload in _sections(d):
        lines += [title.upper(), "-" * len(title)]
        if kind == "text":
            lines.append(payload)
        elif kind == "bullets":
            lines += [f"- {b}" for b in payload]
        else:
            for item in payload:
                t, meta, bullets = _entry(item, kind)
                lines.append(t)
                if meta:
                    lines.append(meta)
                lines += [f"  - {b}" for b in bullets]
        lines.append("")
    return "\n".join(lines)


def build_resume_file(fmt, d, font_key, accent_hex, photo, tpl="classic"):
    """Return (bytes, media_type, filename) for one export format."""
    if fmt == "pdf":
        return build_resume_pdf(d, font_key, accent_hex, photo, tpl), "application/pdf", "formatted_resume.pdf"
    if fmt == "docx":
        return build_resume_docx(d, font_key, accent_hex, photo), MIME_DOCX, "formatted_resume.docx"
    if fmt == "doc":
        docx = build_resume_docx(d, font_key, accent_hex, photo)
        return build_resume_doc(docx), "application/msword", "formatted_resume.doc"
    if fmt == "txt":
        return build_resume_txt(d).encode("utf-8"), "text/plain; charset=utf-8", "formatted_resume.txt"
    raise ValueError("Unknown format. Use DOC, DOCX, PDF, or TXT.")