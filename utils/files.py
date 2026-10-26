import re
import shutil
import subprocess
import tempfile
from html import unescape
from io import BytesIO
from pathlib import Path

SUPPORTED_EXTENSIONS = {".txt", ".html", ".htm", ".rtf", ".pdf", ".docx", ".doc"}
MAX_FILE_BYTES = 5 * 1024 * 1024


def extract_text(filename: str, data: bytes) -> str:
    """Return plain text from an uploaded resume file, or raise ValueError."""
    ext = Path(filename or "").suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError("Unsupported file type. Use DOC, DOCX, PDF, HTML, RTF, or TXT.")
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("File is larger than 5 MB.")

    if ext == ".txt":
        return data.decode("utf-8", errors="replace")

    if ext in (".html", ".htm"):
        html = data.decode("utf-8", errors="replace")
        html = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", html)
        html = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</li>|</h\d>", "\n", html)
        return unescape(re.sub(r"<[^>]+>", " ", html))

    if ext == ".rtf":
        from striprtf.striprtf import rtf_to_text
        return rtf_to_text(data.decode("latin-1", errors="replace"))

    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if ext == ".docx":
        return _docx_text(BytesIO(data))

    return _doc_text(data)


def _docx_text(source) -> str:
    from docx import Document
    doc = Document(source)
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                lines.append(cell.text)
    return "\n".join(lines)


def _doc_text(data: bytes) -> str:
    """Convert a legacy .doc file to DOCX with LibreOffice, then read it."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise ValueError(
            "DOC files need LibreOffice for conversion. "
            "Save the file as DOCX or PDF and upload it again."
        )
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "resume.doc"
        src.write_bytes(data)
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "docx", "--outdir", tmp, str(src)],
                check=True, capture_output=True, timeout=120,
            )
        except (subprocess.SubprocessError, OSError) as e:
            raise ValueError(f"Could not convert DOC file: {e}")
        return _docx_text(Path(tmp) / "resume.docx")